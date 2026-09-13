"""Shared openpyxl helpers for the M7 human-annotation workbooks."""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

EDITABLE_FILL = PatternFill(start_color="FFF6E6", end_color="FFF6E6", fill_type="solid")
HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
HEADER_FONT = Font(bold=True)

LABEL_CHOICES = '"match,non-match,uncertain"'
CONTEXT_CHOICES = '"yes,no"'

WIDE_TEXT_COLS = {"string_a", "string_b", "seed_string", "candidate_string", "justification", "title_1", "title_2", "title_3"}
NARROW_COLS = {"pair_id", "retrieval_pair_id", "domain", "row_number", "label", "context_used", "keyword_string"}


def write_data_sheet(wb: Workbook, sheet_name: str, df: pd.DataFrame, editable_cols: list[str]):
    ws: Worksheet = wb.create_sheet(sheet_name)
    ws.append(list(df.columns))
    for c_idx, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=c_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    for _, row in df.iterrows():
        ws.append([row[c] for c in df.columns])

    # column widths + wrap
    for c_idx, col in enumerate(df.columns, start=1):
        letter = ws.cell(row=1, column=c_idx).column_letter
        width = 40 if col in WIDE_TEXT_COLS else (14 if col in NARROW_COLS else 22)
        ws.column_dimensions[letter].width = width
        if col in WIDE_TEXT_COLS:
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=c_idx).alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # data validation dropdowns
    if "label" in df.columns:
        col_idx = list(df.columns).index("label") + 1
        letter = ws.cell(row=1, column=col_idx).column_letter
        dv = DataValidation(type="list", formula1=LABEL_CHOICES, allow_blank=True, showDropDown=False)
        dv.error = "Choose one of: match, non-match, uncertain"
        dv.prompt = "match / non-match / uncertain"
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{ws.max_row}")
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=col_idx).fill = EDITABLE_FILL

    if "context_used" in df.columns:
        col_idx = list(df.columns).index("context_used") + 1
        letter = ws.cell(row=1, column=col_idx).column_letter
        dv2 = DataValidation(type="list", formula1=CONTEXT_CHOICES, allow_blank=True, showDropDown=False)
        dv2.error = "Choose one of: yes, no"
        dv2.prompt = "yes / no"
        ws.add_data_validation(dv2)
        dv2.add(f"{letter}2:{letter}{ws.max_row}")
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=col_idx).fill = EDITABLE_FILL

    if "justification" in df.columns:
        col_idx = list(df.columns).index("justification") + 1
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=col_idx).fill = EDITABLE_FILL

    # protect non-editable columns (sheet protection with unlocked editable cells)
    ws.protection.sheet = True  # convenience lock (no password), not a security boundary
    for c_idx, col in enumerate(df.columns, start=1):
        editable = col in editable_cols
        for r in range(1, ws.max_row + 1):
            ws.cell(row=r, column=c_idx).protection = ws.cell(row=r, column=c_idx).protection.copy(locked=not editable)
    return ws


def write_instructions_sheet(wb: Workbook, lines: list[str]):
    ws = wb.create_sheet("Instructions", 0)
    ws.column_dimensions["A"].width = 110
    for i, line in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if line.startswith("#"):
            cell.font = Font(bold=True, size=13)
    return ws
