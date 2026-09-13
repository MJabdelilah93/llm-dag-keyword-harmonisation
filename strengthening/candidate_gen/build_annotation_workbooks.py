"""Build the six annotator-facing XLSX workbooks (+ CSV backups) from the
already-derived, gitignored intermediate CSVs. Writes only under
strengthening/restricted_local/human_annotation/v1/.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from .xlsx_helpers import write_data_sheet, write_instructions_sheet

STRENGTHENING_ROOT = Path(__file__).resolve().parents[1]
INTER_DIR = STRENGTHENING_ROOT / "restricted_local" / "human_annotation" / "v1" / "_intermediate"
PKG_DIR = STRENGTHENING_ROOT / "restricted_local" / "human_annotation" / "v1"
CSV_DIR = PKG_DIR / "csv_backups"

PRIMARY_INSTRUCTIONS = [
    "# M7 Primary Benchmark Annotation Instructions",
    "",
    "You are labelling pairs of keyword/concept strings for a scientific benchmark. "
    "Read both strings and decide whether they denote the SAME underlying concept.",
    "",
    "# Allowed labels",
    "",
    "MATCH -- the two strings denote the same underlying concept and can safely be harmonised "
    "(treated as one concept). This includes:",
    "  - spelling variants",
    "  - punctuation/hyphenation variants",
    "  - singular/plural variants (when concept-equivalent)",
    "  - unambiguous acronym / expanded-form pairs",
    "  - genuine synonyms",
    "",
    "NON-MATCH -- the two strings denote different concepts. This includes:",
    "  - broader/narrower relations (one concept is a superset or subset of the other)",
    "  - related but distinct concepts",
    "  - superficially similar terms that in fact mean different things",
    "",
    "UNCERTAIN -- use ONLY when equivalence cannot be determined reliably from the strings "
    "and any permitted bounded context. This includes:",
    "  - ambiguous acronyms (could expand to more than one plausible concept)",
    "  - context-dependent meaning",
    "  - malformed or underspecified expressions",
    "",
    "# What you see and do not see",
    "",
    "You are shown only: a pair (or row) identifier, the domain, the two strings, and blank "
    "label/justification/context_used fields for you to fill in. You are NOT shown any "
    "similarity score, retrieval route, sampling stratum, frequency, or any system/model "
    "prediction -- none exists for these pairs; your judgement is the first and only label.",
    "",
    "# Optional bounded context",
    "",
    "You may consult the accompanying context-lookup file (up to 3 representative article "
    "titles per string) ONLY when you cannot judge the pair from the strings alone. Record "
    "context_used = yes if you consulted it for a given pair, otherwise context_used = no. "
    "Do not search the internet, external databases, or system/model outputs for the answer.",
    "",
    "# Blinding",
    "",
    "Do not discuss individual pair decisions with the other annotator before both of you have "
    "independently completed your full workbook. Two independent annotators complete this same "
    "set of 900 pairs (in different row orders); a third adjudicator resolves any disagreements "
    "afterward.",
    "",
    "# What to fill in",
    "",
    "For every row: set 'label' (match / non-match / uncertain, dropdown provided), 'justification' "
    "(a short free-text reason), and 'context_used' (yes / no, dropdown provided). Leave nothing "
    "blank when you are done. Do not edit pair_id, domain, string_a, or string_b.",
    "",
    "# Returning your work",
    "",
    "Save your completed file using the exact filename given in the accompanying protocol "
    "document (e.g. ANNOTATOR_1_PRIMARY_COMPLETED.xlsx). Do not rename or reorder pair_id values.",
]

RETRIEVAL_INSTRUCTIONS = [
    "# M7 Retrieval-Audit Annotation Instructions",
    "",
    "This is a SEPARATE, secondary validation task -- it checks whether the automated candidate "
    "retrieval process finds real concept-equivalence pairs, and does not feed into or replace "
    "the primary 900-pair benchmark you may also be completing.",
    "",
    "For each row, you see a seed concept and a candidate concept. Decide whether they denote the "
    "same underlying concept using the SAME three labels and definitions as the primary benchmark "
    "(see below) -- match / non-match / uncertain.",
    "",
    "# Allowed labels (same definitions as the primary benchmark)",
    "",
    "MATCH -- same underlying concept (spelling/punctuation/hyphenation/singular-plural variants, "
    "unambiguous acronym/expanded-form pairs, genuine synonyms).",
    "NON-MATCH -- different concepts (broader/narrower, related-but-distinct, superficially similar "
    "but different).",
    "UNCERTAIN -- cannot be determined reliably from the strings and permitted context (ambiguous "
    "acronyms, context-dependent meaning, malformed/underspecified expressions).",
    "",
    "# What you see and do not see",
    "",
    "You see only a retrieval-audit row identifier, domain, the seed string, the candidate string, "
    "and blank label/justification/context_used fields. You do NOT see how the candidate was "
    "retrieved, any similarity score, or whether the row was sampled for any particular reason -- "
    "this information is deliberately withheld to keep your judgement independent of it.",
    "",
    "# Optional bounded context",
    "",
    "As with the primary benchmark, you may consult the accompanying context-lookup file (up to 3 "
    "titles per string) only when needed, and must record context_used = yes/no accordingly.",
    "",
    "# Blinding",
    "",
    "Do not discuss individual row decisions with the other annotator before both have finished. "
    "Annotator 2 works independently and does not see Annotator 1's labels.",
    "",
    "# Returning your work",
    "",
    "Save using the exact filename given in the protocol document (e.g. "
    "ANNOTATOR_1_RETRIEVAL_COMPLETED.xlsx). Do not rename or reorder the identifiers.",
]

PRIMARY_EDITABLE = ["label", "justification", "context_used"]
RETRIEVAL_EDITABLE = ["label", "justification", "context_used"]


def build_primary_workbook(csv_path: Path, out_name: str):
    df = pd.read_csv(csv_path, dtype=str).fillna("")
    ce = df[df["domain"] == "circular_economy"].reset_index(drop=True)
    bio = df[df["domain"] == "biomedical_diabetes_mellitus"].reset_index(drop=True)

    wb = Workbook()
    wb.remove(wb.active)
    write_instructions_sheet(wb, PRIMARY_INSTRUCTIONS)
    write_data_sheet(wb, "Circular_Economy_400", ce, PRIMARY_EDITABLE)
    write_data_sheet(wb, "Diabetes_500", bio, PRIMARY_EDITABLE)
    wb.save(PKG_DIR / out_name)
    print(f"Wrote {PKG_DIR / out_name} (CE={len(ce)}, Diabetes={len(bio)})")


def build_retrieval_workbook(csv_path: Path, out_name: str):
    df = pd.read_csv(csv_path, dtype=str).fillna("")
    ce = df[df["domain"] == "circular_economy"].reset_index(drop=True)
    bio = df[df["domain"] == "biomedical_diabetes_mellitus"].reset_index(drop=True)

    wb = Workbook()
    wb.remove(wb.active)
    write_instructions_sheet(wb, RETRIEVAL_INSTRUCTIONS)
    write_data_sheet(wb, "Circular_Economy_Retrieval", ce, RETRIEVAL_EDITABLE)
    write_data_sheet(wb, "Diabetes_Retrieval", bio, RETRIEVAL_EDITABLE)
    wb.save(PKG_DIR / out_name)
    print(f"Wrote {PKG_DIR / out_name} (CE={len(ce)}, Diabetes={len(bio)})")


def build_context_workbook(csv_path: Path, out_name: str, sheet_name: str):
    df = pd.read_csv(csv_path, dtype=str).fillna("")
    wb = Workbook()
    wb.remove(wb.active)
    write_data_sheet(wb, sheet_name, df, editable_cols=[])
    wb.save(PKG_DIR / out_name)
    print(f"Wrote {PKG_DIR / out_name} ({len(df)} rows)")


def write_csv_backup(csv_path: Path, out_name: str):
    CSV_DIR.mkdir(parents=True, exist_ok=True)
    df = pd.read_csv(csv_path)
    df.to_csv(CSV_DIR / out_name, index=False, encoding="utf-8")


def main():
    PKG_DIR.mkdir(parents=True, exist_ok=True)

    build_primary_workbook(INTER_DIR / "primary_annotator_1.csv", "01_ANNOTATOR_1_PRIMARY.xlsx")
    build_primary_workbook(INTER_DIR / "primary_annotator_2.csv", "02_ANNOTATOR_2_PRIMARY.xlsx")
    build_retrieval_workbook(INTER_DIR / "retrieval_annotator_1.csv", "03_ANNOTATOR_1_RETRIEVAL.xlsx")
    build_retrieval_workbook(INTER_DIR / "retrieval_annotator_2.csv", "04_ANNOTATOR_2_RETRIEVAL.xlsx")
    build_context_workbook(INTER_DIR / "context_lookup_ce.csv", "05_CONTEXT_LOOKUP_CE.xlsx", "CE_Context")
    build_context_workbook(INTER_DIR / "context_lookup_diabetes.csv", "06_CONTEXT_LOOKUP_DIABETES.xlsx", "Diabetes_Context")

    for src, name in [
        (INTER_DIR / "primary_annotator_1.csv", "01_ANNOTATOR_1_PRIMARY.csv"),
        (INTER_DIR / "primary_annotator_2.csv", "02_ANNOTATOR_2_PRIMARY.csv"),
        (INTER_DIR / "retrieval_annotator_1.csv", "03_ANNOTATOR_1_RETRIEVAL.csv"),
        (INTER_DIR / "retrieval_annotator_2.csv", "04_ANNOTATOR_2_RETRIEVAL.csv"),
        (INTER_DIR / "context_lookup_ce.csv", "05_CONTEXT_LOOKUP_CE.csv"),
        (INTER_DIR / "context_lookup_diabetes.csv", "06_CONTEXT_LOOKUP_DIABETES.csv"),
    ]:
        write_csv_backup(src, name)
    print(f"CSV backups written to {CSV_DIR}")


if __name__ == "__main__":
    main()
