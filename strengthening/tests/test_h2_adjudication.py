"""Synthetic-data tests for the H2 pipeline: validation, pair_id-based
merge/agreement, A/B-anonymised adjudication package, the adjudicator
GUI's session logic, and finalise_primary_gold. Never touches real
annotator output."""
from __future__ import annotations

import csv

import pandas as pd
import pytest
from openpyxl import Workbook

from strengthening.human_annotation.finalise_primary_gold import finalise, write_gold_file
from strengthening.human_annotation.h2_agreement import (
    agreement_summary,
    confusion_matrix,
    context_use_crosstab,
    disagreement_types,
)
from strengthening.human_annotation.h2_build_adjudication_package import (
    build_adjudication_dataframe,
    choose_ab_mapping,
)
from strengthening.human_annotation.h2_validate import H2ValidationError, validate_completed
from strengthening.human_annotation.gui.adjudication_session import AdjudicationSession


def _write_completed_workbook(path, rows, shuffle=False):
    """rows: list of (pair_id, domain, string_a, string_b, label, justification, context_used)."""
    if shuffle:
        rows = list(reversed(rows))
    wb = Workbook()
    wb.remove(wb.active)
    ce_rows = [r for r in rows if r[1] == "circular_economy"]
    bio_rows = [r for r in rows if r[1] == "biomedical_diabetes_mellitus"]
    cols = ["pair_id", "domain", "string_a", "string_b", "label", "justification", "context_used"]
    for name, subset in (("Circular_Economy_400", ce_rows), ("Diabetes_500", bio_rows)):
        ws = wb.create_sheet(name)
        ws.append(cols)
        for r in subset:
            ws.append(list(r))
    wb.save(path)


def _canonical(rows):
    return pd.DataFrame(
        [{"pair_id": r[0], "domain": r[1], "string_a": r[2], "string_b": r[3]} for r in rows]
    )


SAMPLE_ROWS = [
    ("ce_1", "circular_economy", "circular economy", "circular  economy", "match", "", "no"),
    ("ce_2", "circular_economy", "recycling", "upcycling", "non-match", "", "no"),
    ("bio_1", "biomedical_diabetes_mellitus", "diabetes mellitus", "diabetic", "match", "", "no"),
    ("bio_2", "biomedical_diabetes_mellitus", "insulin resistance", "insulin sensitivity", "non-match", "", "no"),
]


# -- Step 1: validation ---------------------------------------------------

def test_validate_completed_passes_on_well_formed_workbook(tmp_path):
    p = tmp_path / "ann1.xlsx"
    _write_completed_workbook(p, SAMPLE_ROWS)
    canon = _canonical(SAMPLE_ROWS)
    result = validate_completed(p, canon, "Annotator 1", expected_total=4, expected_ce=2, expected_diabetes=2)
    assert result.ok
    assert result.row_count == 4
    assert result.missing_labels == 0


def test_validate_completed_detects_missing_label(tmp_path):
    rows = list(SAMPLE_ROWS)
    rows[0] = (rows[0][0], rows[0][1], rows[0][2], rows[0][3], "", "", "no")  # blank label
    p = tmp_path / "ann1.xlsx"
    _write_completed_workbook(p, rows)
    canon = _canonical(SAMPLE_ROWS)
    result = validate_completed(p, canon, "Annotator 1", expected_total=4, expected_ce=2, expected_diabetes=2)
    assert not result.ok
    assert result.missing_labels == 1


def test_validate_completed_detects_changed_string(tmp_path):
    rows = list(SAMPLE_ROWS)
    rows[0] = (rows[0][0], rows[0][1], "CHANGED STRING", rows[0][3], rows[0][4], "", "no")
    p = tmp_path / "ann1.xlsx"
    _write_completed_workbook(p, rows)
    canon = _canonical(SAMPLE_ROWS)
    result = validate_completed(p, canon, "Annotator 1", expected_total=4, expected_ce=2, expected_diabetes=2)
    assert not result.ok
    assert "ce_1" in result.changed_strings


def test_validate_completed_detects_invalid_label(tmp_path):
    rows = list(SAMPLE_ROWS)
    rows[0] = (rows[0][0], rows[0][1], rows[0][2], rows[0][3], "definitely-a-match", "", "no")
    p = tmp_path / "ann1.xlsx"
    _write_completed_workbook(p, rows)
    canon = _canonical(SAMPLE_ROWS)
    result = validate_completed(p, canon, "Annotator 1", expected_total=4, expected_ce=2, expected_diabetes=2)
    assert not result.ok
    assert "definitely-a-match" in result.invalid_labels


def test_validate_completed_detects_duplicate_pair_id(tmp_path):
    rows = list(SAMPLE_ROWS) + [SAMPLE_ROWS[0]]
    p = tmp_path / "ann1.xlsx"
    _write_completed_workbook(p, rows)
    canon = _canonical(SAMPLE_ROWS)
    result = validate_completed(p, canon, "Annotator 1", expected_total=4, expected_ce=2, expected_diabetes=2)
    assert not result.ok
    assert "ce_1" in result.duplicate_pair_ids


def test_validate_completed_rejects_missing_file(tmp_path):
    canon = _canonical(SAMPLE_ROWS)
    with pytest.raises(H2ValidationError):
        validate_completed(tmp_path / "does_not_exist.xlsx", canon, "Annotator 1")


# -- Step 2/3: pair_id alignment despite different row order + agreement --

def test_merge_aligns_by_pair_id_not_row_order(tmp_path):
    from strengthening.human_annotation.merge_primary_annotations import merge

    p1 = tmp_path / "ann1.xlsx"
    p2 = tmp_path / "ann2.xlsx"
    _write_completed_workbook(p1, SAMPLE_ROWS, shuffle=False)
    _write_completed_workbook(p2, SAMPLE_ROWS, shuffle=True)  # reversed row order
    merged = merge(p1, p2)
    assert list(merged["pair_id"]) == sorted(merged["pair_id"])
    row = merged[merged["pair_id"] == "ce_1"].iloc[0]
    assert row["annotator_1_label"] == "match" and row["annotator_2_label"] == "match"
    assert row["agree"]


def test_agreement_summary_hand_computed():
    df = pd.DataFrame(
        {
            "annotator_1_label": ["match", "match", "non-match", "uncertain"],
            "annotator_2_label": ["match", "non-match", "non-match", "uncertain"],
        }
    )
    df["agree"] = df["annotator_1_label"] == df["annotator_2_label"]
    s = agreement_summary(df)
    assert s["n"] == 4
    assert s["agreements"] == 3
    assert s["disagreements"] == 1
    assert s["raw_agreement_proportion"] == 0.75


def test_confusion_matrix_hand_computed():
    df = pd.DataFrame({"annotator_1_label": ["match", "non-match"], "annotator_2_label": ["match", "match"]})
    cm = confusion_matrix(df)
    assert cm["match"]["match"] == 1
    assert cm["non-match"]["match"] == 1
    assert cm["match"]["non-match"] == 0


def test_disagreement_types_unordered_bucketing():
    df = pd.DataFrame(
        {
            "annotator_1_label": ["match", "non-match", "match"],
            "annotator_2_label": ["non-match", "match", "uncertain"],
            "agree": [False, False, False],
        }
    )
    types = disagreement_types(df)
    assert types["unordered"]["match vs non-match"] == 2
    assert types["unordered"]["match vs uncertain"] == 1


def test_context_use_crosstab_counts_and_disagreement_rate():
    df = pd.DataFrame(
        {
            "annotator_1_context_used": ["yes", "no", "no", "yes"],
            "annotator_2_context_used": ["no", "no", "yes", "yes"],
            "agree": [False, True, True, False],
        }
    )
    xtab = context_use_crosstab(df)
    assert xtab["neither"]["n"] == 1
    assert xtab["annotator_1_only"]["n"] == 1
    assert xtab["annotator_2_only"]["n"] == 1
    assert xtab["both"]["n"] == 1
    assert xtab["annotator_1_only"]["disagreement_rate"] == 1.0
    assert xtab["neither"]["disagreement_rate"] == 0.0


# -- Step 5: A/B anonymisation ---------------------------------------------

def test_choose_ab_mapping_is_one_of_two_fixed_options():
    m = choose_ab_mapping(rng_seed=1)
    assert m in ({"annotator_1": "A", "annotator_2": "B"}, {"annotator_1": "B", "annotator_2": "A"})


def test_choose_ab_mapping_deterministic_with_seed():
    m1 = choose_ab_mapping(rng_seed=7)
    m2 = choose_ab_mapping(rng_seed=7)
    assert m1 == m2


def test_build_adjudication_dataframe_extracts_disagreements_only_and_anonymises():
    merged = pd.DataFrame(
        {
            "pair_id": ["p1", "p2", "p3"],
            "domain": ["circular_economy"] * 3,
            "string_a": ["a1", "a2", "a3"],
            "string_b": ["b1", "b2", "b3"],
            "annotator_1_label": ["match", "non-match", "match"],
            "annotator_1_justification": ["j1", "j2", "j3"],
            "annotator_1_context_used": ["no", "no", "no"],
            "annotator_2_label": ["match", "match", "non-match"],
            "annotator_2_justification": ["k1", "k2", "k3"],
            "annotator_2_context_used": ["no", "no", "no"],
            "agree": [True, False, False],
        }
    )
    mapping = {"annotator_1": "A", "annotator_2": "B"}
    adj = build_adjudication_dataframe(merged, mapping)
    assert len(adj) == 2  # only p2, p3 disagree
    assert set(adj["pair_id"]) == {"p2", "p3"}
    row_p2 = adj[adj["pair_id"] == "p2"].iloc[0]
    assert row_p2["decision_A"] == "non-match" and row_p2["decision_B"] == "match"
    assert (adj["adjudicated_label"] == "").all()


def test_build_adjudication_dataframe_reverse_mapping_swaps_columns():
    merged = pd.DataFrame(
        {
            "pair_id": ["p1"], "domain": ["circular_economy"], "string_a": ["a"], "string_b": ["b"],
            "annotator_1_label": ["match"], "annotator_1_justification": [""], "annotator_1_context_used": ["no"],
            "annotator_2_label": ["non-match"], "annotator_2_justification": [""], "annotator_2_context_used": ["no"],
            "agree": [False],
        }
    )
    adj_normal = build_adjudication_dataframe(merged, {"annotator_1": "A", "annotator_2": "B"})
    adj_reversed = build_adjudication_dataframe(merged, {"annotator_1": "B", "annotator_2": "A"})
    assert adj_normal.iloc[0]["decision_A"] == adj_reversed.iloc[0]["decision_B"]
    assert adj_normal.iloc[0]["decision_B"] == adj_reversed.iloc[0]["decision_A"]


# -- adjudicator GUI session logic -----------------------------------------

def _write_adjudication_source(path, n=4):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Disagreements")
    cols = ["adjudication_row", "pair_id", "domain", "string_a", "string_b", "decision_A", "decision_B",
            "justification_A", "justification_B", "context_used_A", "context_used_B",
            "adjudicated_label", "adjudicator_notes", "adjudicator_context_used"]
    ws.append(cols)
    for i in range(n):
        ws.append([i + 1, f"p{i}", "circular_economy", f"a{i}", f"b{i}", "match", "non-match", "", "", "no", "no", "", "", ""])
    wb.save(path)


@pytest.fixture
def adj_paths(tmp_path):
    return {
        "source": tmp_path / "PRIMARY_ADJUDICATION.xlsx",
        "working": tmp_path / "PRIMARY_ADJUDICATION_WORKING.xlsx",
        "completed": tmp_path / "PRIMARY_ADJUDICATION_COMPLETED.xlsx",
        "audit": tmp_path / "logs" / "audit.csv",
        "backup": tmp_path / "backups",
    }


def test_adjudicator_autosave_and_resume(adj_paths):
    _write_adjudication_source(adj_paths["source"])
    s1 = AdjudicationSession(adj_paths["source"], adj_paths["working"], adj_paths["completed"], adj_paths["audit"], adj_paths["backup"])
    first_id = s1.current_row().pair_id
    s1.decide("match")

    s2 = AdjudicationSession(adj_paths["source"], adj_paths["working"], adj_paths["completed"], adj_paths["audit"], adj_paths["backup"])
    assert s2._by_id[first_id].adjudicated_label == "match"


def test_adjudicator_skip_leaves_blank_and_back_relabels(adj_paths):
    _write_adjudication_source(adj_paths["source"])
    s = AdjudicationSession(adj_paths["source"], adj_paths["working"], adj_paths["completed"], adj_paths["audit"], adj_paths["backup"])
    first_id = s.current_row().pair_id
    s.skip()
    assert s._by_id[first_id].adjudicated_label == ""
    s.go_back()
    assert s.current_row().pair_id == first_id
    s.decide("uncertain")
    assert s._by_id[first_id].adjudicated_label == "uncertain"
    # relabel
    s.go_back()
    s.decide("match")
    assert s._by_id[first_id].adjudicated_label == "match"


def test_adjudicator_context_logging(adj_paths):
    _write_adjudication_source(adj_paths["source"])
    s = AdjudicationSession(adj_paths["source"], adj_paths["working"], adj_paths["completed"], adj_paths["audit"], adj_paths["backup"])
    s.mark_context_opened()
    assert s.current_row().adjudicator_context_used == "yes"
    s.decide("match")
    assert s._by_id[s.rows[0].pair_id].adjudicator_context_used == "yes"  # preserved, not overwritten to "no"

    # a different, never-context-opened row defaults to "no" upon decision
    s.jump_to(s.rows[1].pair_id)
    assert s.current_row().adjudicator_context_used == ""
    s.decide("non-match")
    assert s._by_id[s.rows[1].pair_id].adjudicator_context_used == "no"


def test_adjudicator_audit_log_records_old_and_new(adj_paths):
    _write_adjudication_source(adj_paths["source"])
    s = AdjudicationSession(adj_paths["source"], adj_paths["working"], adj_paths["completed"], adj_paths["audit"], adj_paths["backup"])
    first_id = s.current_row().pair_id
    s.decide("match")
    s.go_back()
    s.decide("non-match")
    rows = list(csv.DictReader(open(adj_paths["audit"], newline="", encoding="utf-8")))
    matching = [r for r in rows if r["pair_id"] == first_id]
    assert matching[0]["old_label"] == "" and matching[0]["new_label"] == "match"
    assert matching[1]["old_label"] == "match" and matching[1]["new_label"] == "non-match" and matching[1]["action"] == "relabel"


def test_adjudicator_completed_output_only_when_all_labelled(adj_paths):
    _write_adjudication_source(adj_paths["source"], n=2)
    s = AdjudicationSession(adj_paths["source"], adj_paths["working"], adj_paths["completed"], adj_paths["audit"], adj_paths["backup"])
    s.decide("match")
    assert s.write_completed_if_done() is None  # only 1/2 done
    s.decide("non-match")
    out = s.write_completed_if_done()
    assert out == adj_paths["completed"]
    assert out.exists()


def test_adjudicator_never_overwrites_existing_completed_file(adj_paths):
    _write_adjudication_source(adj_paths["source"], n=1)
    s = AdjudicationSession(adj_paths["source"], adj_paths["working"], adj_paths["completed"], adj_paths["audit"], adj_paths["backup"])
    s.decide("match")
    out1 = s.write_completed_if_done()
    before = out1.read_bytes()
    # simulate calling it again (e.g. app re-render) -- must not silently rewrite
    out2 = s.write_completed_if_done()
    assert out2.read_bytes() == before


# -- Step 7: finalise_primary_gold (synthetic only) ------------------------

def test_finalise_primary_gold_keeps_agreed_and_replaces_disagreed():
    merged_900 = pd.DataFrame(
        {
            "pair_id": ["p1", "p2"],
            "domain": ["circular_economy", "circular_economy"],
            "string_a": ["a1", "a2"],
            "string_b": ["b1", "b2"],
            "annotator_1_label": ["match", "non-match"],
            "annotator_1_justification": ["", ""],
            "annotator_1_context_used": ["no", "no"],
            "annotator_2_label": ["match", "uncertain"],
            "annotator_2_justification": ["", ""],
            "annotator_2_context_used": ["no", "no"],
            "agree": [True, False],
        }
    )
    adjudication_completed = pd.DataFrame(
        {"pair_id": ["p2"], "adjudicated_label": ["uncertain"], "adjudicator_notes": ["resolved"], "adjudicator_context_used": ["no"]}
    )
    gold = finalise(merged_900, adjudication_completed)
    p1 = gold[gold["pair_id"] == "p1"].iloc[0]
    p2 = gold[gold["pair_id"] == "p2"].iloc[0]
    assert p1["final_gold_label"] == "match" and p1["gold_source"] == "direct_agreement"
    assert p2["final_gold_label"] == "uncertain" and p2["gold_source"] == "adjudicated"
    assert p2["annotator_1_label"] == "non-match" and p2["annotator_2_corrected_label"] == "uncertain"  # both preserved


def test_finalise_primary_gold_raises_on_incomplete_adjudication():
    merged_900 = pd.DataFrame(
        {
            "pair_id": ["p1"], "domain": ["circular_economy"], "string_a": ["a"], "string_b": ["b"],
            "annotator_1_label": ["match"], "annotator_1_justification": [""], "annotator_1_context_used": ["no"],
            "annotator_2_label": ["non-match"], "annotator_2_justification": [""], "annotator_2_context_used": ["no"],
            "agree": [False],
        }
    )
    adjudication_incomplete = pd.DataFrame({"pair_id": ["p1"], "adjudicated_label": [""], "adjudicator_notes": [""], "adjudicator_context_used": [""]})
    with pytest.raises(ValueError):
        finalise(merged_900, adjudication_incomplete)


def test_write_gold_file_refuses_to_overwrite_and_checks_row_count(tmp_path):
    gold = pd.DataFrame({"pair_id": ["p1"]})
    out = tmp_path / "gold.csv"
    with pytest.raises(ValueError):
        write_gold_file(gold, out)  # not 900 rows

    gold_900 = pd.DataFrame({"pair_id": [f"p{i}" for i in range(900)]})
    write_gold_file(gold_900, out)
    with pytest.raises(FileExistsError):
        write_gold_file(gold_900, out)  # already exists


# -- restricted files gitignored -------------------------------------------

def test_h2_restricted_directory_is_gitignored():
    import subprocess
    from pathlib import Path

    worktree = Path(__file__).resolve().parents[2]
    if not (worktree / ".git").exists():
        pytest.skip("requires a git checkout to verify gitignore coverage via `git check-ignore`; not applicable to a plain source archive/tarball extraction")
    proc = subprocess.run(
        ["git", "check-ignore", "-v", "strengthening/restricted_local/human_annotation/v1/h2/PRIMARY_ADJUDICATION.xlsx"],
        cwd=worktree, capture_output=True, text=True,
    )
    assert proc.returncode == 0


def test_adjudicator_bat_uses_the_fixed_launcher_pattern():
    from pathlib import Path

    import pytest

    # As of the diabetes-annotation-quality diagnostic, the LIVE
    # START_PRIMARY_ADJUDICATOR.bat is intentionally a blocking stub (the
    # 478-pair adjudication package is superseded pending diabetes
    # re-annotation -- see ADJUDICATION_STATUS.txt). The original,
    # functional launcher content is preserved unchanged under the
    # ".DISABLED_..." filename; that is what this test now checks.
    h2_dir = Path(__file__).resolve().parents[1] / "restricted_local" / "human_annotation" / "v1" / "h2"
    disabled = h2_dir / "START_PRIMARY_ADJUDICATOR.bat.DISABLED_PENDING_DIABETES_REANNOTATION"
    if not disabled.exists():
        pytest.skip("disabled launcher backup not present in this environment")
    text = disabled.read_text(encoding="utf-8")
    assert 'if "%SCRIPT_DIR:~-1%"=="\\" set "SCRIPT_DIR=%SCRIPT_DIR:~0,-1%"' in text
    assert 'start "" /min python.exe' in text
    assert "goto :launch_pythonw" in text


def test_adjudication_gui_module_exposes_no_forbidden_metadata():
    import ast
    import inspect

    from strengthening.human_annotation.gui import adjudication_app

    # Scan actual code, not the module's own docstring (which legitimately
    # documents, in prose, which fields must never be shown -- that mention
    # is the safety property being tested for, not a violation of it).
    source = inspect.getsource(adjudication_app)
    tree = ast.parse(source)
    module_docstring = ast.get_docstring(tree) or ""
    code_only = source.replace(module_docstring, "", 1).lower()
    for term in ("stratum", "jaro", "tfidf", "embedding", "frequency", "route", "claude", "anthropic", "openai", "confidence", "gold", "guard"):
        assert term not in code_only, f"forbidden term '{term}' found outside the module docstring"
