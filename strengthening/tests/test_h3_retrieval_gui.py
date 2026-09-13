"""Tests for the H3 retrieval-audit GUI (session/app/entrypoint) and the
future per-row-anonymised retrieval adjudication scaffold. Synthetic data
only, except hash/schema/launcher-pattern checks against the real,
never-modified frozen retrieval package on disk."""
from __future__ import annotations

import hashlib
import subprocess
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from strengthening.human_annotation.build_retrieval_adjudication_package import (
    build_anonymised,
    choose_ab_mapping_per_row,
)
from strengthening.human_annotation.gui.retrieval_main import (
    EXPECTED_ROWS,
    FORBIDDEN_COLUMN_TERMS,
    REQUIRED_RETRIEVAL_COLUMNS,
    validate_retrieval_source,
)
from strengthening.human_annotation.gui.retrieval_session import (
    ALLOWED_LABELS,
    RetrievalAnnotationSession,
)
from strengthening.human_annotation.gui.validation import ValidationError

PKG_DIR = Path(__file__).resolve().parents[1] / "restricted_local" / "human_annotation" / "v1"
INTERMEDIATE_DIR = PKG_DIR / "_intermediate"


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _write_retrieval_workbook(path, ce_rows, bio_rows):
    wb = Workbook()
    wb.remove(wb.active)
    cols = ["retrieval_pair_id", "domain", "seed_string", "candidate_string", "label", "justification", "context_used"]
    for name, subset in (("Circular_Economy_Retrieval", ce_rows), ("Diabetes_Retrieval", bio_rows)):
        ws = wb.create_sheet(name)
        ws.append(cols)
        for r in subset:
            ws.append(list(r))
    wb.save(path)


SAMPLE_CE = [
    ("r_ce_1", "circular_economy", "seed one", "candidate one", "", "", ""),
    ("r_ce_2", "circular_economy", "seed two", "candidate two", "", "", ""),
]
SAMPLE_BIO = [
    ("r_bio_1", "biomedical_diabetes_mellitus", "seed three", "candidate three", "", "", ""),
]


# -- retrieval source validation ---------------------------------------------

def test_validate_retrieval_source_passes_on_well_formed_workbook(tmp_path, monkeypatch):
    p = tmp_path / "retrieval.xlsx"
    _write_retrieval_workbook(p, SAMPLE_CE, SAMPLE_BIO)
    monkeypatch.setitem(EXPECTED_ROWS, "1", 3)
    validate_retrieval_source(p, "1")  # must not raise


def test_validate_retrieval_source_rejects_wrong_row_count(tmp_path, monkeypatch):
    p = tmp_path / "retrieval.xlsx"
    _write_retrieval_workbook(p, SAMPLE_CE, SAMPLE_BIO)
    monkeypatch.setitem(EXPECTED_ROWS, "1", 999)
    with pytest.raises(ValidationError):
        validate_retrieval_source(p, "1")


def test_validate_retrieval_source_rejects_forbidden_column(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Circular_Economy_Retrieval")
    cols = ["retrieval_pair_id", "domain", "seed_string", "candidate_string", "label", "justification", "context_used", "embedding_cosine"]
    ws.append(cols)
    ws.append(["r1", "circular_economy", "a", "b", "", "", "", "0.8"])
    p = tmp_path / "leaky.xlsx"
    wb.save(p)
    with pytest.raises(ValidationError):
        validate_retrieval_source(p, "1")


def test_validate_retrieval_source_rejects_duplicate_ids(tmp_path, monkeypatch):
    p = tmp_path / "retrieval.xlsx"
    _write_retrieval_workbook(p, SAMPLE_CE + [SAMPLE_CE[0]], SAMPLE_BIO)
    monkeypatch.setitem(EXPECTED_ROWS, "1", 4)
    with pytest.raises(ValidationError):
        validate_retrieval_source(p, "1")


def test_validate_retrieval_source_missing_file_raises(tmp_path):
    with pytest.raises(ValidationError):
        validate_retrieval_source(tmp_path / "nope.xlsx", "1")


def test_required_retrieval_columns_match_frozen_schema():
    assert REQUIRED_RETRIEVAL_COLUMNS == ["retrieval_pair_id", "domain", "seed_string", "candidate_string", "label", "justification", "context_used"]


def test_forbidden_column_terms_cover_all_blinding_categories():
    for term in ("stratum", "jaro", "embedding", "route", "rank", "confidence", "gold", "outside_pool", "audit_sample"):
        assert term in FORBIDDEN_COLUMN_TERMS


# -- label policy preserved ---------------------------------------------------

def test_allowed_labels_unchanged():
    assert ALLOWED_LABELS == ("match", "non-match", "uncertain")


# -- session: autosave / resume / skip / back-relabel / completion ----------

@pytest.fixture
def retrieval_paths(tmp_path):
    return {
        "source": tmp_path / "source.xlsx",
        "working": tmp_path / "WORKING.xlsx",
        "completed": tmp_path / "COMPLETED.xlsx",
        "audit": tmp_path / "logs" / "audit.csv",
        "backup": tmp_path / "backups",
    }


def test_retrieval_session_autosave_and_resume(retrieval_paths):
    _write_retrieval_workbook(retrieval_paths["source"], SAMPLE_CE, SAMPLE_BIO)
    s1 = RetrievalAnnotationSession("A1", retrieval_paths["source"], retrieval_paths["working"], retrieval_paths["completed"], retrieval_paths["audit"], retrieval_paths["backup"])
    first_id = s1.current_pair().retrieval_pair_id
    s1.decide("match")

    s2 = RetrievalAnnotationSession("A1", retrieval_paths["source"], retrieval_paths["working"], retrieval_paths["completed"], retrieval_paths["audit"], retrieval_paths["backup"])
    assert s2._by_id[first_id].label == "match"
    assert s2.current_pair().retrieval_pair_id != first_id  # resumed at first UNRESOLVED row


def test_retrieval_session_skip_leaves_blank_and_back_relabels(retrieval_paths):
    _write_retrieval_workbook(retrieval_paths["source"], SAMPLE_CE, SAMPLE_BIO)
    s = RetrievalAnnotationSession("A1", retrieval_paths["source"], retrieval_paths["working"], retrieval_paths["completed"], retrieval_paths["audit"], retrieval_paths["backup"])
    first_id = s.current_pair().retrieval_pair_id
    s.skip()
    assert s._by_id[first_id].label == ""
    s.go_back()
    assert s.current_pair().retrieval_pair_id == first_id
    s.decide("uncertain")
    assert s._by_id[first_id].label == "uncertain"
    s.go_back()
    s.decide("match")  # relabel
    assert s._by_id[first_id].label == "match"


def test_retrieval_session_context_use_defaults_and_preserved(retrieval_paths):
    _write_retrieval_workbook(retrieval_paths["source"], SAMPLE_CE, SAMPLE_BIO)
    s = RetrievalAnnotationSession("A1", retrieval_paths["source"], retrieval_paths["working"], retrieval_paths["completed"], retrieval_paths["audit"], retrieval_paths["backup"])
    s.mark_context_opened()
    assert s.current_pair().context_used == "yes"
    s.decide("match")
    assert s._by_id[s.pairs[0].retrieval_pair_id].context_used == "yes"

    s.jump_to(s.pairs[1].retrieval_pair_id)
    assert s.current_pair().context_used == ""
    s.decide("non-match")
    assert s._by_id[s.pairs[1].retrieval_pair_id].context_used == "no"


def test_retrieval_completed_file_only_after_all_rows_labelled(retrieval_paths):
    _write_retrieval_workbook(retrieval_paths["source"], SAMPLE_CE[:1], [])
    s = RetrievalAnnotationSession("A1", retrieval_paths["source"], retrieval_paths["working"], retrieval_paths["completed"], retrieval_paths["audit"], retrieval_paths["backup"])
    assert s.write_completed_if_done() is None
    s.decide("match")
    out = s.write_completed_if_done()
    assert out == retrieval_paths["completed"] and out.exists()


def test_retrieval_completed_file_never_overwritten(retrieval_paths):
    _write_retrieval_workbook(retrieval_paths["source"], SAMPLE_CE[:1], [])
    s = RetrievalAnnotationSession("A1", retrieval_paths["source"], retrieval_paths["working"], retrieval_paths["completed"], retrieval_paths["audit"], retrieval_paths["backup"])
    s.decide("match")
    out1 = s.write_completed_if_done()
    before = out1.read_bytes()
    out2 = s.write_completed_if_done()
    assert out2.read_bytes() == before


def test_retrieval_session_backs_up_every_100_decisions(retrieval_paths):
    rows = [(f"r{i}", "circular_economy", f"s{i}", f"c{i}", "", "", "") for i in range(150)]
    _write_retrieval_workbook(retrieval_paths["source"], rows, [])
    s = RetrievalAnnotationSession("A1", retrieval_paths["source"], retrieval_paths["working"], retrieval_paths["completed"], retrieval_paths["audit"], retrieval_paths["backup"])
    for _ in range(100):
        s.decide("match")
    assert retrieval_paths["backup"].exists()
    assert len(list(retrieval_paths["backup"].glob("*.xlsx"))) == 1


def test_retrieval_session_rejects_invalid_label(retrieval_paths):
    _write_retrieval_workbook(retrieval_paths["source"], SAMPLE_CE, SAMPLE_BIO)
    s = RetrievalAnnotationSession("A1", retrieval_paths["source"], retrieval_paths["working"], retrieval_paths["completed"], retrieval_paths["audit"], retrieval_paths["backup"])
    with pytest.raises(ValueError):
        s.decide("definitely-a-match")


def test_retrieval_session_detects_drift_between_source_and_working(retrieval_paths):
    _write_retrieval_workbook(retrieval_paths["source"], SAMPLE_CE, SAMPLE_BIO)
    RetrievalAnnotationSession("A1", retrieval_paths["source"], retrieval_paths["working"], retrieval_paths["completed"], retrieval_paths["audit"], retrieval_paths["backup"])
    # simulate a corrupted/wrong-annotator working file: change a seed string
    tampered_rows = [(SAMPLE_CE[0][0], SAMPLE_CE[0][1], "TAMPERED", SAMPLE_CE[0][3], "", "", "")] + SAMPLE_CE[1:]
    _write_retrieval_workbook(retrieval_paths["working"], tampered_rows, SAMPLE_BIO)
    with pytest.raises(ValueError):
        RetrievalAnnotationSession("A1", retrieval_paths["source"], retrieval_paths["working"], retrieval_paths["completed"], retrieval_paths["audit"], retrieval_paths["backup"])


# -- future per-row-anonymised retrieval adjudication scaffold ---------------

def test_choose_ab_mapping_per_row_deterministic_and_balanced():
    ids = [f"r{i}" for i in range(1000)]
    m1 = choose_ab_mapping_per_row(ids, seed=11)
    m2 = choose_ab_mapping_per_row(ids, seed=11)
    assert m1 == m2
    n_a = sum(1 for m in m1.values() if m["annotator_1"] == "A")
    assert 0.4 < n_a / len(m1) < 0.6


def test_build_anonymised_retrieval_package_pairs_fields_and_blank_labels():
    merged = pd.DataFrame({
        "retrieval_pair_id": ["r1"], "domain": ["circular_economy"], "seed_string": ["s1"], "candidate_string": ["c1"],
        "annotator_1_label": ["match"], "annotator_1_justification": ["j1"],
        "annotator_2_label": ["non-match"], "annotator_2_justification": ["j2"],
        "double_coded": [True], "agree": [False],
    })
    mapping = {"r1": {"annotator_1": "A", "annotator_2": "B"}}
    pkg = build_anonymised(merged, mapping)
    assert len(pkg) == 1
    assert pkg.iloc[0]["decision_A"] == "match" and pkg.iloc[0]["decision_B"] == "non-match"
    assert (pkg["adjudicated_label"] == "").all()
    assert "annotator_1_label" not in pkg.columns and "annotator_2_label" not in pkg.columns


# -- real (never-modified) frozen H3 package ---------------------------------

def test_frozen_retrieval_workbooks_unchanged():
    expected = {
        "03_ANNOTATOR_1_RETRIEVAL.xlsx": "404204cac1274a6e5ae70a50939f6099843c4993e1a00171552a84ec68d25fae",
        "04_ANNOTATOR_2_RETRIEVAL.xlsx": "6cbefd6237df762e4cabb067aac4909cd3f2e9cfacf8919a5a6e154d8fe40dae",
    }
    for name, expected_hash in expected.items():
        path = PKG_DIR / name
        if not path.exists():
            pytest.skip(f"{name} not present in this environment")
        assert _sha256(path) == expected_hash


def test_frozen_retrieval_workbook_real_row_counts():
    path = PKG_DIR / "03_ANNOTATOR_1_RETRIEVAL.xlsx"
    if not path.exists():
        pytest.skip("real retrieval package not present in this environment")
    sheets = pd.read_excel(path, sheet_name=None, dtype=str)
    ce = sheets["Circular_Economy_Retrieval"]
    bio = sheets["Diabetes_Retrieval"]
    assert len(ce) == 3117 and len(bio) == 3068
    assert len(ce) + len(bio) == 6185
    assert ce["label"].notna().sum() == 0 and bio["label"].notna().sum() == 0

    a2_path = PKG_DIR / "04_ANNOTATOR_2_RETRIEVAL.xlsx"
    a2_sheets = pd.read_excel(a2_path, sheet_name=None, dtype=str)
    a2_ce, a2_bio = a2_sheets["Circular_Economy_Retrieval"], a2_sheets["Diabetes_Retrieval"]
    assert len(a2_ce) == 1039 and len(a2_bio) == 1027
    assert len(a2_ce) + len(a2_bio) == 2066
    assert set(a2_ce["retrieval_pair_id"]).issubset(set(ce["retrieval_pair_id"]))
    assert set(a2_bio["retrieval_pair_id"]).issubset(set(bio["retrieval_pair_id"]))


def test_frozen_a2_subset_matches_outside_pool_plus_audit_sample():
    ce_master_path = INTERMEDIATE_DIR / "scenario_e_ce_master.csv"
    bio_master_path = INTERMEDIATE_DIR / "scenario_e_bio_master.csv"
    a2_path = PKG_DIR / "04_ANNOTATOR_2_RETRIEVAL.xlsx"
    if not (ce_master_path.exists() and bio_master_path.exists() and a2_path.exists()):
        pytest.skip("intermediate master files or real A2 package not present in this environment")
    ce_master = pd.read_csv(ce_master_path)
    bio_master = pd.read_csv(bio_master_path)
    a2_sheets = pd.read_excel(a2_path, sheet_name=None, dtype=str)

    ce_expected = set(ce_master.loc[ce_master["outside_pool_sample"] | ce_master["audit_sample_selected"], "retrieval_pair_id"])
    bio_expected = set(bio_master.loc[bio_master["outside_pool_sample"] | bio_master["audit_sample_selected"], "retrieval_pair_id"])
    assert set(a2_sheets["Circular_Economy_Retrieval"]["retrieval_pair_id"]) == ce_expected
    assert set(a2_sheets["Diabetes_Retrieval"]["retrieval_pair_id"]) == bio_expected
    assert ce_master["outside_pool_sample"].sum() == 150 and ce_master["audit_sample_selected"].sum() == 889
    assert bio_master["outside_pool_sample"].sum() == 150 and bio_master["audit_sample_selected"].sum() == 877


def test_no_working_or_completed_retrieval_files_exist():
    matches = list(PKG_DIR.glob("ANNOTATOR_*_RETRIEVAL_WORKING*")) + list(PKG_DIR.glob("ANNOTATOR_*_RETRIEVAL_COMPLETED*"))
    assert matches == [], f"unexpected retrieval progress files found: {matches}"


# -- launcher pattern (Windows path-with-spaces / trailing-backslash) -------

def test_both_retrieval_launchers_use_the_fixed_pattern_and_own_isolated_source():
    launcher_1 = PKG_DIR / "START_ANNOTATOR_1_RETRIEVAL.bat"
    launcher_2 = PKG_DIR / "START_ANNOTATOR_2_RETRIEVAL.bat"
    if not (launcher_1.exists() and launcher_2.exists()):
        pytest.skip("retrieval launchers not present in this environment")
    text_1 = launcher_1.read_text(encoding="utf-8")
    text_2 = launcher_2.read_text(encoding="utf-8")

    for text in (text_1, text_2):
        assert 'if "%SCRIPT_DIR:~-1%"=="\\" set "SCRIPT_DIR=%SCRIPT_DIR:~0,-1%"' in text
        assert 'start "" /min python.exe' in text
        assert "goto :launch_pythonw" in text
        assert "retrieval_main" in text

    # annotator isolation: each launcher references ONLY its own source file
    assert "03_ANNOTATOR_1_RETRIEVAL.xlsx" in text_1 and "04_ANNOTATOR_2_RETRIEVAL.xlsx" not in text_1
    assert "04_ANNOTATOR_2_RETRIEVAL.xlsx" in text_2 and "03_ANNOTATOR_1_RETRIEVAL.xlsx" not in text_2
    assert "--annotator-id 1" in text_1 and "--annotator-id 2" in text_2


def test_retrieval_launcher_argv_survives_a_path_containing_spaces(tmp_path):
    """Empirically reproduces the exact %~dp0 pattern used by the retrieval
    launchers inside a directory that has a space in its name (like the
    real worktree path), and inspects the REAL child process's sys.argv
    via a stub script -- proving the fixed pattern still avoids the
    trailing-backslash-before-quote bug for these new launchers."""
    space_dir = tmp_path / "Article 7 space test"
    space_dir.mkdir()
    stub = space_dir / "print_argv.py"
    stub.write_text("import sys\nprint('ARGV:' + '|'.join(sys.argv[1:]))\n", encoding="utf-8")

    bat_path = space_dir / "test_launcher.bat"
    bat_path.write_text(
        '@echo off\r\n'
        'setlocal\r\n'
        'set "SCRIPT_DIR=%~dp0"\r\n'
        'if "%SCRIPT_DIR:~-1%"=="\\" set "SCRIPT_DIR=%SCRIPT_DIR:~0,-1%"\r\n'
        'python.exe "%SCRIPT_DIR%\\print_argv.py" --source "%SCRIPT_DIR%\\03_ANNOTATOR_1_RETRIEVAL.xlsx" --package-dir "%SCRIPT_DIR%"\r\n',
        encoding="utf-8",
    )
    proc = subprocess.run(["cmd.exe", "/c", str(bat_path)], capture_output=True, text=True, cwd=str(space_dir))
    out_line = next(line for line in proc.stdout.splitlines() if line.startswith("ARGV:"))
    args = out_line[len("ARGV:"):].split("|")
    assert args[0] == "--source"
    assert args[1] == str(space_dir / "03_ANNOTATOR_1_RETRIEVAL.xlsx")
    assert args[2] == "--package-dir"
    assert args[3] == str(space_dir)  # no leaked literal quote from a trailing backslash


def test_retrieval_gitignore_coverage():
    worktree = Path(__file__).resolve().parents[2]
    if not (worktree / ".git").exists():
        pytest.skip("requires a git checkout to verify gitignore coverage via `git check-ignore`; not applicable to a plain source archive/tarball extraction")
    proc = subprocess.run(
        ["git", "check-ignore", "-v", "strengthening/restricted_local/human_annotation/v1/ANNOTATOR_1_RETRIEVAL_WORKING.xlsx"],
        cwd=worktree, capture_output=True, text=True,
    )
    assert proc.returncode == 0
