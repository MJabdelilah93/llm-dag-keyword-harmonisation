"""Gold-freeze Step 2: independent validation of the COMPLETED corrected
adjudication workbook. Checked against the CORRECTED H2 disagreement set
(never against the superseded 478-row package) and against the pristine
pre-adjudication source, to detect any drift in fields the adjudicator
must never be able to change (string_a/string_b, domain, decision_A/B,
justifications, context-use metadata). Fails loudly rather than silently
repairing anything.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ALLOWED_LABELS = {"match", "non-match", "uncertain"}
ALLOWED_CONTEXT_USED = {"yes", "no"}

REQUIRED_COLS = [
    "adjudication_row", "pair_id", "domain", "string_a", "string_b",
    "decision_A", "decision_B", "justification_A", "justification_B",
    "context_used_A", "context_used_B",
    "adjudicated_label", "adjudicator_notes", "adjudicator_context_used",
]
IMMUTABLE_COLS = [
    "domain", "string_a", "string_b", "decision_A", "decision_B",
    "justification_A", "justification_B", "context_used_A", "context_used_B",
]
FORBIDDEN_COLUMN_TERMS = [
    "stratum", "jaro", "tfidf", "tf_idf", "embedding", "frequency", "route",
    "rank", "confidence", "gold", "guard", "claude", "openai", "anthropic", "model",
]


class AdjudicationValidationError(Exception):
    pass


@dataclass
class AdjudicationValidationResult:
    ok: bool
    row_count: int = 0
    duplicate_pair_ids: list = field(default_factory=list)
    missing_pair_ids: list = field(default_factory=list)
    extra_pair_ids: list = field(default_factory=list)
    blank_adjudicated_labels: int = 0
    invalid_labels: list = field(default_factory=list)
    invalid_context_used: list = field(default_factory=list)
    drifted_pair_ids: list = field(default_factory=list)
    forbidden_columns: list = field(default_factory=list)
    issues: list = field(default_factory=list)


def _read_sheet(path: Path) -> pd.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    rows = [dict(zip(header, values)) for values in it]
    wb.close()
    return pd.DataFrame(rows)


def validate_completed_adjudication(
    completed_path: Path,
    source_path: Path,
    expected_disagreement_ids: set,
    expected_rows: int = 286,
) -> AdjudicationValidationResult:
    if not completed_path.exists():
        raise AdjudicationValidationError(f"Completed adjudication workbook not found: {completed_path}")
    if not source_path.exists():
        raise AdjudicationValidationError(f"Pristine pre-adjudication source not found: {source_path}")

    completed = _read_sheet(completed_path)
    source = _read_sheet(source_path)

    result = AdjudicationValidationResult(ok=True)

    missing_cols = [c for c in REQUIRED_COLS if c not in completed.columns]
    if missing_cols:
        raise AdjudicationValidationError(f"Completed adjudication workbook missing required columns: {missing_cols}")

    extra_cols = [c for c in completed.columns if c not in REQUIRED_COLS]
    forbidden = [c for c in extra_cols if any(t in c.lower() for t in FORBIDDEN_COLUMN_TERMS)]
    if forbidden:
        result.forbidden_columns = forbidden
        result.ok = False
        result.issues.append(f"completed workbook exposes forbidden columns: {forbidden}")

    result.row_count = len(completed)
    if result.row_count != expected_rows:
        result.ok = False
        result.issues.append(f"expected {expected_rows} rows, found {result.row_count}")

    dup = completed["pair_id"][completed["pair_id"].duplicated()]
    if len(dup):
        result.duplicate_pair_ids = sorted(set(dup))
        result.ok = False
        result.issues.append(f"{len(dup)} duplicated pair_id rows")

    completed_ids = set(completed["pair_id"])
    missing = expected_disagreement_ids - completed_ids
    extra = completed_ids - expected_disagreement_ids
    if missing:
        result.missing_pair_ids = sorted(missing)[:5]
        result.ok = False
        result.issues.append(f"{len(missing)} corrected-H2 disagreement pair_ids missing from completed adjudication")
    if extra:
        result.extra_pair_ids = sorted(extra)[:5]
        result.ok = False
        result.issues.append(f"{len(extra)} pair_ids in completed adjudication are not part of the corrected-H2 disagreement set")

    blank_mask = completed["adjudicated_label"].isna() | (completed["adjudicated_label"].astype(str).str.strip() == "")
    result.blank_adjudicated_labels = int(blank_mask.sum())
    if result.blank_adjudicated_labels:
        result.ok = False
        result.issues.append(f"{result.blank_adjudicated_labels} rows have a blank adjudicated_label")

    bad_label_mask = ~completed["adjudicated_label"].isin(ALLOWED_LABELS) & ~blank_mask
    if bad_label_mask.any():
        result.invalid_labels = sorted(set(completed.loc[bad_label_mask, "adjudicated_label"]))
        result.ok = False
        result.issues.append(f"invalid adjudicated_label values: {result.invalid_labels}")

    bad_ctx_mask = ~completed["adjudicator_context_used"].isin(ALLOWED_CONTEXT_USED)
    if bad_ctx_mask.any():
        result.invalid_context_used = sorted(set(completed.loc[bad_ctx_mask, "adjudicator_context_used"].astype(str)))
        result.ok = False
        result.issues.append(f"invalid/blank adjudicator_context_used values: {result.invalid_context_used}")

    # drift check: only adjudicated_label/adjudicator_notes/adjudicator_context_used
    # may differ from the pristine pre-adjudication source.
    source_by_id = source.set_index("pair_id")
    completed_by_id = completed.drop_duplicates(subset="pair_id", keep="first").set_index("pair_id")
    drifted = []
    for pid in completed_by_id.index:
        if pid not in source_by_id.index:
            continue
        c, s = completed_by_id.loc[pid], source_by_id.loc[pid]
        if tuple(c[col] for col in IMMUTABLE_COLS) != tuple(s[col] for col in IMMUTABLE_COLS):
            drifted.append(pid)
    if drifted:
        result.drifted_pair_ids = drifted
        result.ok = False
        result.issues.append(
            f"{len(drifted)} pair_ids have drifted immutable fields (decision_A/B, justifications, "
            f"context-use metadata, string_a/string_b, domain) vs the pristine pre-adjudication source"
        )

    return result


if __name__ == "__main__":
    raise SystemExit(
        "This module validates the REAL completed corrected adjudication workbook. Invoke it via "
        "strengthening/human_annotation/build_primary_gold_freeze.py, not directly."
    )
