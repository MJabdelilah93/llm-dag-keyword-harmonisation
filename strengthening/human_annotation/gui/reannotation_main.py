"""M7 diabetes-only re-annotation GUI entrypoint. Invoked by
START_DIABETES_REANNOTATION_ANNOTATOR_2.bat, never directly. The source
workbook is a BRAND-NEW, blank-label, freshly-ordered file that never
contains the original Annotator 2 diabetes labels, Annotator 1 labels, or
any model/system information -- independence from prior annotation is
structural (the source data simply does not have it), not a GUI filter.
"""
from __future__ import annotations

import argparse
import sys
import tkinter as tk
from pathlib import Path
from tkinter import messagebox

from .context_lookup import CombinedContextLookup
from .main import clean_path_arg
from .reannotation_app import ReannotationApp
from .session import AnnotationSession
from .validation import ValidationError, validate_context_lookup


def _fatal_error_dialog(message: str) -> None:
    root = tk.Tk()
    root.withdraw()
    messagebox.showerror("M7 Diabetes Re-annotation - Cannot Start", message)
    root.destroy()


def validate_reannotation_source(source: Path) -> None:
    """Dedicated, lightweight validator for the 500-row diabetes-only
    re-annotation source (deliberately NOT the 900/400/500-row primary
    validate_source_workbook, which does not apply here and is left
    untouched to avoid any regression risk to the already-tested primary
    GUI)."""
    from openpyxl import load_workbook

    if not source.exists():
        raise ValidationError(f"Re-annotation source workbook does not exist: {source}")
    wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    idx = {c: i for i, c in enumerate(header)}
    required = ["pair_id", "domain", "string_a", "string_b", "label", "justification", "context_used"]
    missing = [c for c in required if c not in idx]
    if missing:
        wb.close()
        raise ValidationError(f"Re-annotation source missing required columns: {missing}")

    rows = list(it)
    wb.close()
    if len(rows) != 500:
        raise ValidationError(f"Expected exactly 500 diabetes re-annotation rows, found {len(rows)}")
    pair_ids = [r[idx["pair_id"]] for r in rows]
    if len(set(pair_ids)) != 500:
        raise ValidationError("Re-annotation source has duplicate or missing pair_id values")
    domains = {r[idx["domain"]] for r in rows}
    if domains != {"biomedical_diabetes_mellitus"}:
        raise ValidationError(f"Re-annotation source must be diabetes-only, found domains: {domains}")
    non_blank_labels = [r[idx["label"]] for r in rows if r[idx["label"]]]
    if non_blank_labels:
        raise ValidationError(
            f"Re-annotation source must start with every label blank, found {len(non_blank_labels)} pre-filled labels"
        )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="M7 diabetes-only re-annotation GUI")
    parser.add_argument("--source", required=True, type=clean_path_arg)
    parser.add_argument("--context-ce", required=True, type=clean_path_arg)
    parser.add_argument("--context-diabetes", required=True, type=clean_path_arg)
    parser.add_argument("--package-dir", required=True, type=clean_path_arg)
    args = parser.parse_args(argv)

    try:
        validate_reannotation_source(args.source)
        validate_context_lookup(args.context_ce)
        validate_context_lookup(args.context_diabetes)
    except ValidationError as e:
        _fatal_error_dialog(str(e))
        return 1
    except Exception as e:  # noqa: BLE001
        _fatal_error_dialog(f"Unexpected error while validating inputs:\n{e}")
        return 1

    package_dir: Path = args.package_dir
    working_path = package_dir / "DIABETES_REANNOTATION_ANNOTATOR_2_WORKING.xlsx"
    completed_path = package_dir / "DIABETES_REANNOTATION_ANNOTATOR_2_COMPLETED.xlsx"
    audit_log_path = package_dir / "logs" / "DIABETES_REANNOTATION_ANNOTATOR_2_audit_log.csv"
    backup_dir = package_dir / "backups"

    try:
        session = AnnotationSession(
            annotator_id="Annotator 2 (diabetes re-annotation)",
            source_path=args.source,
            working_path=working_path,
            completed_path=completed_path,
            audit_log_path=audit_log_path,
            backup_dir=backup_dir,
        )
        context = CombinedContextLookup(args.context_ce, args.context_diabetes)
    except Exception as e:  # noqa: BLE001
        _fatal_error_dialog(f"Could not start the re-annotation session:\n{e}")
        return 1

    app = ReannotationApp(session, context, "Annotator 2 (diabetes re-annotation)")
    app.run()
    return 0


if __name__ == "__main__":
    sys.exit(main())
