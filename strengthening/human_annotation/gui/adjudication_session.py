"""Core adjudication-session state machine: load/resume/autosave, back/
skip/decide navigation, audit logging, and completed-file production.
Mirrors session.py's design (kept separate from Tkinter) but operates on
disagreement rows (decision_A/decision_B already known, adjudicated_label
is the one field being decided) rather than a fresh two-way label.
"""
from __future__ import annotations

import csv
import os
import shutil
import tempfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

ALLOWED_LABELS = ("match", "non-match", "uncertain")
DATA_TAB_NAME = "Disagreements"
COLUMNS = [
    "adjudication_row", "pair_id", "domain", "string_a", "string_b",
    "decision_A", "decision_B", "justification_A", "justification_B",
    "context_used_A", "context_used_B",
    "adjudicated_label", "adjudicator_notes", "adjudicator_context_used",
]


@dataclass
class DisagreementRow:
    adjudication_row: int
    pair_id: str
    domain: str
    string_a: str
    string_b: str
    decision_A: str
    decision_B: str
    justification_A: str
    justification_B: str
    context_used_A: str
    context_used_B: str
    adjudicated_label: str = ""
    adjudicator_notes: str = ""
    adjudicator_context_used: str = ""

    def as_row(self) -> list:
        return [getattr(self, c) for c in COLUMNS]


def _s(v) -> str:
    return v if isinstance(v, str) else ("" if v is None else str(v))


def _read_rows(xlsx_path: Path) -> list[DisagreementRow]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    idx = {c: i for i, c in enumerate(header)}
    rows = []
    for values in it:
        rows.append(
            DisagreementRow(
                adjudication_row=values[idx["adjudication_row"]],
                pair_id=values[idx["pair_id"]],
                domain=values[idx["domain"]],
                string_a=values[idx["string_a"]],
                string_b=values[idx["string_b"]],
                decision_A=values[idx["decision_A"]],
                decision_B=values[idx["decision_B"]],
                justification_A=_s(values[idx["justification_A"]]),
                justification_B=_s(values[idx["justification_B"]]),
                context_used_A=_s(values[idx["context_used_A"]]),
                context_used_B=_s(values[idx["context_used_B"]]),
                adjudicated_label=_s(values[idx["adjudicated_label"]]) if "adjudicated_label" in idx else "",
                adjudicator_notes=_s(values[idx["adjudicator_notes"]]) if "adjudicator_notes" in idx else "",
                adjudicator_context_used=_s(values[idx["adjudicator_context_used"]]) if "adjudicator_context_used" in idx else "",
            )
        )
    wb.close()
    return rows


def save_workbook(rows: list[DisagreementRow], path: Path) -> None:
    """Atomic save: temp file in the same directory + os.replace."""
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(DATA_TAB_NAME)
    ws.append(COLUMNS)
    for r in rows:
        ws.append(r.as_row())

    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(dir=str(path.parent), prefix=".tmp_save_", suffix=".xlsx")
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink(missing_ok=True)


class AdjudicationSession:
    BACKUP_EVERY_N_DECISIONS = 25

    def __init__(self, source_path: Path, working_path: Path, completed_path: Path, audit_log_path: Path, backup_dir: Path):
        self.source_path = Path(source_path)
        self.working_path = Path(working_path)
        self.completed_path = Path(completed_path)
        self.audit_log_path = Path(audit_log_path)
        self.backup_dir = Path(backup_dir)

        self._source_rows = _read_rows(self.source_path)
        self._source_by_id = {r.pair_id: r for r in self._source_rows}

        if self.working_path.exists():
            self.rows = _read_rows(self.working_path)
            self._verify_no_drift()
        else:
            self.rows = [
                DisagreementRow(
                    s.adjudication_row, s.pair_id, s.domain, s.string_a, s.string_b,
                    s.decision_A, s.decision_B, s.justification_A, s.justification_B,
                    s.context_used_A, s.context_used_B,
                )
                for s in self._source_rows
            ]
            save_workbook(self.rows, self.working_path)

        self._by_id = {r.pair_id: r for r in self.rows}
        self.current_index = self._first_unresolved_index_or_zero()
        self._decisions_since_backup = 0

    def _verify_no_drift(self) -> None:
        for r in self.rows:
            src = self._source_by_id.get(r.pair_id)
            if src is None:
                raise ValueError(f"Working file has unknown pair_id '{r.pair_id}' not in the pristine adjudication package")
            if (r.string_a, r.string_b, r.domain, r.decision_A, r.decision_B) != (src.string_a, src.string_b, src.domain, src.decision_A, src.decision_B):
                raise ValueError(f"pair_id '{r.pair_id}': content drift between source and working adjudication file")

    def _first_unresolved_index_or_zero(self) -> int:
        for i, r in enumerate(self.rows):
            if not r.adjudicated_label:
                return i
        return 0

    def current_row(self) -> DisagreementRow:
        return self.rows[self.current_index]

    def progress(self) -> tuple[int, int]:
        completed = sum(1 for r in self.rows if r.adjudicated_label)
        return completed, len(self.rows)

    def is_complete(self) -> bool:
        return all(r.adjudicated_label for r in self.rows)

    def _next_unresolved_from(self, start: int) -> int | None:
        n = len(self.rows)
        for offset in range(1, n + 1):
            i = (start + offset) % n
            if not self.rows[i].adjudicated_label:
                return i
        return None

    def go_next_unresolved(self) -> int | None:
        nxt = self._next_unresolved_from(self.current_index)
        if nxt is not None:
            self.current_index = nxt
        return nxt

    def go_back(self) -> int:
        self.current_index = (self.current_index - 1) % len(self.rows)
        return self.current_index

    def jump_to(self, pair_id: str) -> None:
        for i, r in enumerate(self.rows):
            if r.pair_id == pair_id:
                self.current_index = i
                return
        raise KeyError(pair_id)

    def _autosave(self) -> None:
        save_workbook(self.rows, self.working_path)
        self._decisions_since_backup += 1
        if self._decisions_since_backup >= self.BACKUP_EVERY_N_DECISIONS:
            self.backup_now()
            self._decisions_since_backup = 0

    def _log(self, pair_id: str, old_label: str, new_label: str, context_used: str, action: str) -> None:
        is_new = not self.audit_log_path.exists()
        self.audit_log_path.parent.mkdir(parents=True, exist_ok=True)
        with open(self.audit_log_path, "a", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            if is_new:
                writer.writerow(["timestamp", "pair_id", "old_label", "new_label", "context_used", "action"])
            writer.writerow([datetime.now(timezone.utc).isoformat(), pair_id, old_label, new_label, context_used, action])

    def decide(self, label: str) -> None:
        if label not in ALLOWED_LABELS:
            raise ValueError(f"label must be one of {ALLOWED_LABELS}, got {label!r}")
        row = self.current_row()
        old_label = row.adjudicated_label
        row.adjudicated_label = label
        if not row.adjudicator_context_used:
            row.adjudicator_context_used = "no"
        action = "relabel" if old_label and old_label != label else "label"
        self._log(row.pair_id, old_label, label, row.adjudicator_context_used, action)
        self._autosave()
        self.go_next_unresolved()

    def skip(self) -> None:
        row = self.current_row()
        self._log(row.pair_id, row.adjudicated_label, row.adjudicated_label, row.adjudicator_context_used, "skip")
        self._autosave()
        self.go_next_unresolved()

    def set_notes(self, text: str) -> None:
        row = self.current_row()
        row.adjudicator_notes = text
        self._log(row.pair_id, row.adjudicated_label, row.adjudicated_label, row.adjudicator_context_used, "notes_edit")
        self._autosave()

    def mark_context_opened(self) -> None:
        row = self.current_row()
        if row.adjudicator_context_used == "yes":
            return
        row.adjudicator_context_used = "yes"
        self._log(row.pair_id, row.adjudicated_label, row.adjudicated_label, row.adjudicator_context_used, "context_opened")
        self._autosave()

    def backup_now(self) -> Path:
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S")
        dest = self.backup_dir / f"{self.working_path.stem}_backup_{ts}.xlsx"
        shutil.copy2(self.working_path, dest)
        return dest

    def write_completed_if_done(self) -> Path | None:
        if not self.is_complete():
            return None
        if self.completed_path.exists():
            return self.completed_path  # never overwrite an already-finalised completed file
        save_workbook(self.rows, self.completed_path)
        return self.completed_path
