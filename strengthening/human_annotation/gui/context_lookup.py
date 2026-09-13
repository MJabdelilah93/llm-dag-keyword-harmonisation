"""Loads the frozen context-lookup files (up to 3 titles per string).
Never exposes abstracts, keywords, or scores -- the source files
themselves only ever contain keyword_string/title_1/title_2/title_3."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


class ContextLookup:
    def __init__(self, xlsx_path: Path):
        self._titles: dict[str, list[str]] = {}
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = list(next(rows))
        idx = {c: i for i, c in enumerate(header)}
        for values in rows:
            kw = values[idx["keyword_string"]]
            titles = [values[idx[c]] for c in ("title_1", "title_2", "title_3") if c in idx]
            titles = [t for t in titles if t and str(t).strip()]
            self._titles[kw] = titles
        wb.close()

    def titles_for(self, keyword: str) -> list[str]:
        return self._titles.get(keyword, [])


class CombinedContextLookup:
    """Merges the CE and diabetes context lookups (a string only ever
    appears in one of the two domains' files in practice, but merging
    keeps the GUI's call sites domain-agnostic)."""

    def __init__(self, ce_path: Path, diabetes_path: Path):
        self._ce = ContextLookup(ce_path)
        self._diabetes = ContextLookup(diabetes_path)

    def titles_for(self, keyword: str) -> list[str]:
        return self._ce.titles_for(keyword) or self._diabetes.titles_for(keyword)
