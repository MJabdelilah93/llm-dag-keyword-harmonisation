"""Core annotation-session state machine: load/resume/autosave, back/skip/
decide navigation, audit logging, and completed-file production. No
Tkinter here -- this module is unit-testable headlessly.
"""
from __future__ import annotations

import csv
import os
import shutil
import tempfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

ALLOWED_LABELS = ("match", "non-match", "uncertain")
DATA_TAB_NAMES = {"circular_economy": "Circular_Economy_400", "biomedical_diabetes_mellitus": "Diabetes_500"}
COLUMNS = ["pair_id", "domain", "string_a", "string_b", "label", "justification", "context_used"]


@dataclass
class PairState:
    pair_id: str
    domain: str
    string_a: str
    string_b: str
    label: str = ""
    justification: str = ""
    context_used: str = ""

    def as_row(self) -> list:
        return [self.pair_id, self.domain, self.string_a, self.string_b, self.label, self.justification, self.context_used]


@dataclass
class AuditEntry:
    timestamp: str
    annotator_id: str
    pair_id: str
    old_label: str
    new_label: str
    context_used: str
    action: str


def _read_source_rows(xlsx_path: Path) -> list[PairState]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    pairs: list[PairState] = []
    for sheet_name in wb.sheetnames:
        if sheet_name == "Instructions":
            continue
        ws = wb[sheet_name]
        it = ws.iter_rows(values_only=True)
        header = list(next(it))
        idx = {c: i for i, c in enumerate(header)}
        for values in it:
            pairs.append(
                PairState(
                    pair_id=values[idx["pair_id"]],
                    domain=values[idx["domain"]],
                    string_a=values[idx["string_a"]],
                    string_b=values[idx["string_b"]],
                    label=(values[idx["label"]] or "") if "label" in idx else "",
                    justification=(values[idx["justification"]] or "") if "justification" in idx else "",
                    context_used=(values[idx["context_used"]] or "") if "context_used" in idx else "",
                )
            )
    wb.close()
    return pairs


def save_workbook(pairs: list[PairState], path: Path) -> None:
    """Atomic save: write to a temp file in the same directory, then
    os.replace -- an interrupted write can never corrupt the target."""
    wb = Workbook()
    wb.remove(wb.active)
    by_domain: dict[str, list[PairState]] = {}
    for p in pairs:
        by_domain.setdefault(p.domain, []).append(p)
    for domain, tab_name in DATA_TAB_NAMES.items():
        ws = wb.create_sheet(tab_name)
        ws.append(COLUMNS)
        for p in by_domain.get(domain, []):
            ws.append(p.as_row())

    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(dir=str(path.parent), prefix=".tmp_save_", suffix=".xlsx")
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink(missing_ok=True)


class AnnotationSession:
    BACKUP_EVERY_N_DECISIONS = 25

    def __init__(self, annotator_id: str, source_path: Path, working_path: Path, completed_path: Path, audit_log_path: Path, backup_dir: Path):
        self.annotator_id = annotator_id
        self.source_path = Path(source_path)
        self.working_path = Path(working_path)
        self.completed_path = Path(completed_path)
        self.audit_log_path = Path(audit_log_path)
        self.backup_dir = Path(backup_dir)

        self._source_pairs = _read_source_rows(self.source_path)
        self._source_by_id = {p.pair_id: p for p in self._source_pairs}

        if self.working_path.exists():
            self.pairs = _read_source_rows(self.working_path)
            self._verify_no_drift()
        else:
            self.pairs = [PairState(p.pair_id, p.domain, p.string_a, p.string_b) for p in self._source_pairs]
            save_workbook(self.pairs, self.working_path)

        self._by_id = {p.pair_id: p for p in self.pairs}
        self.current_index = self._first_unresolved_index_or_zero()
        self._decisions_since_backup = 0

    def _verify_no_drift(self) -> None:
        for p in self.pairs:
            src = self._source_by_id.get(p.pair_id)
            if src is None:
                raise ValueError(f"Working file has unknown pair_id '{p.pair_id}' not in pristine source")
            if (p.string_a, p.string_b, p.domain) != (src.string_a, src.string_b, src.domain):
                raise ValueError(f"pair_id '{p.pair_id}': string/domain drift between source and working file")

    # -- navigation -----------------------------------------------------
    def _first_unresolved_index_or_zero(self) -> int:
        for i, p in enumerate(self.pairs):
            if not p.label:
                return i
        return 0

    def current_pair(self) -> PairState:
        return self.pairs[self.current_index]

    def progress(self) -> tuple[int, int]:
        completed = sum(1 for p in self.pairs if p.label)
        return completed, len(self.pairs)

    def is_complete(self) -> bool:
        return all(p.label for p in self.pairs)

    def _next_unresolved_from(self, start: int) -> int | None:
        n = len(self.pairs)
        for offset in range(1, n + 1):
            i = (start + offset) % n
            if not self.pairs[i].label:
                return i
        return None

    def go_next_unresolved(self) -> int | None:
        nxt = self._next_unresolved_from(self.current_index)
        if nxt is not None:
            self.current_index = nxt
        return nxt

    def go_back(self) -> int:
        self.current_index = (self.current_index - 1) % len(self.pairs)
        return self.current_index

    def jump_to(self, pair_id: str) -> None:
        for i, p in enumerate(self.pairs):
            if p.pair_id == pair_id:
                self.current_index = i
                return
        raise KeyError(pair_id)

    # -- mutating actions (each autosaves + logs) ------------------------
    def _autosave(self) -> None:
        save_workbook(self.pairs, self.working_path)
        self._decisions_since_backup += 1
        if self._decisions_since_backup >= self.BACKUP_EVERY_N_DECISIONS:
            self.backup_now()
            self._decisions_since_backup = 0

    def _log(self, pair_id: str, old_label: str, new_label: str, context_used: str, action: str) -> None:
        is_new = not self.audit_log_path.exists()
        self.audit_log_path.parent.mkdir(parents=True, exist_ok=True)
        with open(self.audit_log_path, "a", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            if is_new:
                writer.writerow(["timestamp", "annotator_id", "pair_id", "old_label", "new_label", "context_used", "action"])
            writer.writerow([datetime.now(timezone.utc).isoformat(), self.annotator_id, pair_id, old_label, new_label, context_used, action])

    def decide(self, label: str) -> None:
        if label not in ALLOWED_LABELS:
            raise ValueError(f"label must be one of {ALLOWED_LABELS}, got {label!r}")
        pair = self.current_pair()
        old_label = pair.label
        pair.label = label
        if not pair.context_used:
            pair.context_used = "no"
        action = "relabel" if old_label and old_label != label else "label"
        self._log(pair.pair_id, old_label, label, pair.context_used, action)
        self._autosave()
        self.go_next_unresolved()

    def skip(self) -> None:
        pair = self.current_pair()
        self._log(pair.pair_id, pair.label, pair.label, pair.context_used, "skip")
        self._autosave()
        self.go_next_unresolved()

    def set_justification(self, text: str) -> None:
        pair = self.current_pair()
        pair.justification = text
        self._log(pair.pair_id, pair.label, pair.label, pair.context_used, "justification_edit")
        self._autosave()

    def mark_context_opened(self) -> None:
        pair = self.current_pair()
        if pair.context_used == "yes":
            return  # already recorded, avoid noisy duplicate log entries
        pair.context_used = "yes"
        self._log(pair.pair_id, pair.label, pair.label, pair.context_used, "context_opened")
        self._autosave()

    # -- backup / completion ---------------------------------------------
    def backup_now(self) -> Path:
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S")
        dest = self.backup_dir / f"{self.working_path.stem}_backup_{ts}.xlsx"
        shutil.copy2(self.working_path, dest)
        return dest

    def write_completed_if_done(self) -> Path | None:
        if not self.is_complete():
            return None
        save_workbook(self.pairs, self.completed_path)
        return self.completed_path
