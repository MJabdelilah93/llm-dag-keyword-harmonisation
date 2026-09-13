"""M7 H3 retrieval-audit GUI entrypoint. Invoked by
START_ANNOTATOR_<N>_RETRIEVAL.bat, never by the annotator directly.

Usage:
    python -m strengthening.human_annotation.gui.retrieval_main \
        --annotator-id 1 \
        --source "<...>/03_ANNOTATOR_1_RETRIEVAL.xlsx" \
        --context-ce "<...>/05_CONTEXT_LOOKUP_CE.xlsx" \
        --context-diabetes "<...>/06_CONTEXT_LOOKUP_DIABETES.xlsx" \
        --package-dir "<...>/restricted_local/human_annotation/v1"

Each annotator's launcher points --source at ONLY that annotator's own
pristine retrieval workbook (03_... for Annotator 1, 04_... for Annotator
2) -- there is structurally no code path here that could open the other
annotator's file, so isolation does not depend on the annotator behaving
correctly.
"""
from __future__ import annotations

import argparse
import sys
import tkinter as tk
from pathlib import Path
from tkinter import messagebox

from .context_lookup import CombinedContextLookup
from .main import clean_path_arg
from .retrieval_app import RetrievalApp
from .retrieval_session import RetrievalAnnotationSession
from .validation import ValidationError, validate_context_lookup

REQUIRED_RETRIEVAL_COLUMNS = ["retrieval_pair_id", "domain", "seed_string", "candidate_string", "label", "justification", "context_used"]
EXPECTED_ROWS = {"1": 6185, "2": 2066}
FORBIDDEN_COLUMN_TERMS = [
    "stratum", "jaro", "tfidf", "tf_idf", "embedding", "frequency", "route", "rank",
    "confidence", "gold", "guard", "claude", "openai", "anthropic", "model",
    "outside_pool", "audit_sample", "in_pool", "difficulty_band",
]


def _fatal_error_dialog(message: str) -> None:
    root = tk.Tk()
    root.withdraw()
    messagebox.showerror("M7 Retrieval Audit GUI - Cannot Start", message)
    root.destroy()


def validate_retrieval_source(source: Path, annotator_id: str) -> None:
    """Dedicated validator for the retrieval-audit workbook schema
    (deliberately NOT the primary GUI's validate_source_workbook, which
    checks the unrelated 900/400/500 primary schema)."""
    from openpyxl import load_workbook

    if not source.exists():
        raise ValidationError(f"Retrieval workbook does not exist: {source}")
    wb = load_workbook(source, read_only=True, data_only=True)
    rows: list[dict] = []
    for sheet_name in wb.sheetnames:
        if sheet_name == "Instructions":
            continue
        ws = wb[sheet_name]
        it = ws.iter_rows(values_only=True)
        header = list(next(it))
        idx = {c: i for i, c in enumerate(header)}
        missing = [c for c in REQUIRED_RETRIEVAL_COLUMNS if c not in idx]
        if missing:
            wb.close()
            raise ValidationError(f"Retrieval workbook sheet '{sheet_name}' missing required columns: {missing}")
        forbidden = [c for c in header if any(t in str(c).lower() for t in FORBIDDEN_COLUMN_TERMS)]
        if forbidden:
            wb.close()
            raise ValidationError(f"Retrieval workbook exposes forbidden columns: {forbidden}")
        for values in it:
            rows.append(dict(zip(header, values)))
    wb.close()

    expected = EXPECTED_ROWS.get(annotator_id)
    if expected is not None and len(rows) != expected:
        raise ValidationError(f"Expected exactly {expected} retrieval rows for Annotator {annotator_id}, found {len(rows)}")

    ids = [r["retrieval_pair_id"] for r in rows]
    if len(set(ids)) != len(ids):
        raise ValidationError("Retrieval workbook has duplicate or missing retrieval_pair_id values")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="M7 H3 retrieval-audit GUI")
    parser.add_argument("--annotator-id", required=True, choices=["1", "2"])
    parser.add_argument("--source", required=True, type=clean_path_arg)
    parser.add_argument("--context-ce", required=True, type=clean_path_arg)
    parser.add_argument("--context-diabetes", required=True, type=clean_path_arg)
    parser.add_argument("--package-dir", required=True, type=clean_path_arg)
    args = parser.parse_args(argv)

    annotator_label = f"Annotator {args.annotator_id} (retrieval audit)"

    try:
        validate_retrieval_source(args.source, args.annotator_id)
        validate_context_lookup(args.context_ce)
        validate_context_lookup(args.context_diabetes)
    except ValidationError as e:
        _fatal_error_dialog(str(e))
        return 1
    except Exception as e:  # noqa: BLE001
        _fatal_error_dialog(f"Unexpected error while validating inputs:\n{e}")
        return 1

    package_dir: Path = args.package_dir
    working_path = package_dir / f"ANNOTATOR_{args.annotator_id}_RETRIEVAL_WORKING.xlsx"
    completed_path = package_dir / f"ANNOTATOR_{args.annotator_id}_RETRIEVAL_COMPLETED.xlsx"
    audit_log_path = package_dir / "logs" / f"ANNOTATOR_{args.annotator_id}_RETRIEVAL_audit_log.csv"
    backup_dir = package_dir / "backups"

    try:
        session = RetrievalAnnotationSession(
            annotator_id=annotator_label,
            source_path=args.source,
            working_path=working_path,
            completed_path=completed_path,
            audit_log_path=audit_log_path,
            backup_dir=backup_dir,
        )
        context = CombinedContextLookup(args.context_ce, args.context_diabetes)
    except Exception as e:  # noqa: BLE001
        _fatal_error_dialog(f"Could not start the retrieval-audit session:\n{e}")
        return 1

    app = RetrievalApp(session, context, annotator_label)
    app.run()
    return 0


if __name__ == "__main__":
    sys.exit(main())
