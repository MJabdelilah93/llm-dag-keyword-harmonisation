"""M7 H2 primary adjudication GUI entrypoint. Invoked by
START_PRIMARY_ADJUDICATOR.bat, never by the adjudicator directly.

Usage:
    python -m strengthening.human_annotation.gui.adjudication_main \
        --source "<...>/h2/PRIMARY_ADJUDICATION.xlsx" \
        --context-ce "<...>/05_CONTEXT_LOOKUP_CE.xlsx" \
        --context-diabetes "<...>/06_CONTEXT_LOOKUP_DIABETES.xlsx" \
        --h2-dir "<...>/h2"
"""
from __future__ import annotations

import argparse
import sys
import tkinter as tk
from pathlib import Path
from tkinter import messagebox

from .adjudication_app import AdjudicationApp
from .adjudication_session import AdjudicationSession
from .context_lookup import CombinedContextLookup
from .main import clean_path_arg  # the fixed, quote-stripping path type (commit 2f17f49)


def _fatal_error_dialog(message: str) -> None:
    root = tk.Tk()
    root.withdraw()
    messagebox.showerror("M7 Adjudication GUI - Cannot Start", message)
    root.destroy()


def _resolve_paths(h2_dir: Path, filename_prefix: str) -> tuple[Path, Path, Path]:
    """Pure path-construction, factored out so it is unit-testable without
    starting a Tkinter mainloop. Default prefix ("PRIMARY_ADJUDICATION")
    reproduces the original filenames exactly; the corrected package passes
    "PRIMARY_ADJUDICATION_CORRECTED" to get its own working/completed/audit
    files without colliding with the original (superseded) package."""
    working_path = h2_dir / f"{filename_prefix}_WORKING.xlsx"
    completed_path = h2_dir / f"{filename_prefix}_COMPLETED.xlsx"
    audit_log_path = h2_dir / "logs" / f"{filename_prefix}_audit_log.csv"
    return working_path, completed_path, audit_log_path


def _validate_source(source: Path) -> None:
    from openpyxl import load_workbook

    if not source.exists():
        raise FileNotFoundError(f"Adjudication package not found: {source}")
    wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = next(ws.iter_rows(values_only=True))
    required = ["adjudication_row", "pair_id", "domain", "string_a", "string_b", "decision_A", "decision_B", "adjudicated_label"]
    missing = [c for c in required if c not in header]
    wb.close()
    if missing:
        raise ValueError(f"Adjudication package missing required columns: {missing}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="M7 primary adjudication GUI")
    parser.add_argument("--source", required=True, type=clean_path_arg)
    parser.add_argument("--context-ce", required=True, type=clean_path_arg)
    parser.add_argument("--context-diabetes", required=True, type=clean_path_arg)
    parser.add_argument("--h2-dir", required=True, type=clean_path_arg)
    parser.add_argument("--filename-prefix", default="PRIMARY_ADJUDICATION")
    args = parser.parse_args(argv)

    try:
        _validate_source(args.source)
    except Exception as e:  # noqa: BLE001
        _fatal_error_dialog(f"Cannot start adjudication:\n{e}")
        return 1

    h2_dir: Path = args.h2_dir
    working_path, completed_path, audit_log_path = _resolve_paths(h2_dir, args.filename_prefix)
    backup_dir = h2_dir / "backups"

    try:
        session = AdjudicationSession(args.source, working_path, completed_path, audit_log_path, backup_dir)
        context = CombinedContextLookup(args.context_ce, args.context_diabetes)
    except Exception as e:  # noqa: BLE001
        _fatal_error_dialog(f"Could not start the adjudication session:\n{e}")
        return 1

    app = AdjudicationApp(session, context)
    app.run()
    return 0


if __name__ == "__main__":
    sys.exit(main())
