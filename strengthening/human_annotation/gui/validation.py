"""Pre-launch validation for the M7 annotation GUI. Fails loudly (raises
ValidationError) rather than silently continuing, per protocol."""
from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

REQUIRED_PRIMARY_COLUMNS = ["pair_id", "domain", "string_a", "string_b", "label", "justification", "context_used"]
EXPECTED_TOTAL = 900
EXPECTED_CE = 400
EXPECTED_DIABETES = 500


class ValidationError(Exception):
    pass


@dataclass
class ValidationResult:
    ok: bool
    messages: list[str] = field(default_factory=list)
    row_count: int = 0
    ce_count: int = 0
    diabetes_count: int = 0
    source_sha256: str = ""


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _read_data_rows(xlsx_path: Path) -> list[dict]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    rows: list[dict] = []
    for sheet_name in wb.sheetnames:
        if sheet_name == "Instructions":
            continue
        ws = wb[sheet_name]
        it = ws.iter_rows(values_only=True)
        try:
            header = next(it)
        except StopIteration:
            continue
        header = list(header)
        for values in it:
            rows.append(dict(zip(header, values)))
    wb.close()
    return rows


def validate_source_workbook(xlsx_path: Path, expected_manifest_hash: str | None = None) -> ValidationResult:
    messages: list[str] = []
    if not xlsx_path.exists():
        raise ValidationError(f"Source workbook does not exist: {xlsx_path}")

    rows = _read_data_rows(xlsx_path)
    missing_cols = [c for c in REQUIRED_PRIMARY_COLUMNS if rows and c not in rows[0]]
    if missing_cols:
        raise ValidationError(f"Source workbook missing required columns: {missing_cols}")

    pair_ids = [r.get("pair_id") for r in rows]
    if len(pair_ids) != len(set(pair_ids)):
        dupes = {p for p in pair_ids if pair_ids.count(p) > 1}
        raise ValidationError(f"Duplicate pair_id values found in source workbook: {sorted(dupes)[:10]}")

    if len(rows) != EXPECTED_TOTAL:
        raise ValidationError(f"Expected exactly {EXPECTED_TOTAL} rows, found {len(rows)}")

    ce_count = sum(1 for r in rows if r.get("domain") == "circular_economy")
    diab_count = sum(1 for r in rows if r.get("domain") == "biomedical_diabetes_mellitus")
    if ce_count != EXPECTED_CE:
        raise ValidationError(f"Expected {EXPECTED_CE} circular_economy rows, found {ce_count}")
    if diab_count != EXPECTED_DIABETES:
        raise ValidationError(f"Expected {EXPECTED_DIABETES} biomedical_diabetes_mellitus rows, found {diab_count}")

    digest = sha256_of(xlsx_path)
    if expected_manifest_hash and digest != expected_manifest_hash:
        messages.append(
            f"WARNING: source workbook sha256 ({digest[:12]}...) does not match the value recorded in the "
            f"annotation-package manifest ({expected_manifest_hash[:12]}...) -- the file may have been "
            "regenerated or modified since packaging."
        )

    return ValidationResult(ok=True, messages=messages, row_count=len(rows), ce_count=ce_count, diabetes_count=diab_count, source_sha256=digest)


def validate_context_lookup(xlsx_path: Path) -> bool:
    if not xlsx_path.exists():
        raise ValidationError(f"Context lookup file does not exist: {xlsx_path}")
    rows = _read_data_rows(xlsx_path)
    if not rows or "keyword_string" not in rows[0]:
        raise ValidationError(f"Context lookup file malformed (missing 'keyword_string' column): {xlsx_path}")
    return True


def validate_no_string_drift(source_rows: list[dict], working_rows: list[dict]) -> None:
    """Used on resume: pair_id -> (string_a, string_b) must match exactly
    between the pristine source and the working file (catches a corrupted
    or wrong-annotator working file)."""
    source_by_id = {r["pair_id"]: (r["string_a"], r["string_b"], r["domain"]) for r in source_rows}
    for r in working_rows:
        pid = r["pair_id"]
        if pid not in source_by_id:
            raise ValidationError(f"Working file has pair_id '{pid}' not present in the pristine source workbook")
        sa, sb, dom = source_by_id[pid]
        if (r["string_a"], r["string_b"], r["domain"]) != (sa, sb, dom):
            raise ValidationError(f"pair_id '{pid}': string_a/string_b/domain differ between source and working file")
